import asyncio
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfWriter

from app.main import app, get_repository
from app.models import CandidateProfile
from app.services.repository import LocalJsonRepository
from app.services.exporter import jobs_to_csv, jobs_to_xlsx
from test_pipeline import make_job, student


@pytest.fixture
def client(tmp_path):
    repo = LocalJsonRepository(tmp_path / "jobs.json")
    asyncio.run(repo.upsert_jobs([make_job(url=f"https://example.com/jobs/{n}",title=f"Frontend Intern {n}") for n in range(15)]))
    asyncio.run(repo.put_state("profile", student().model_dump()))
    app.dependency_overrides[get_repository] = lambda: repo
    with TestClient(app) as client:
        yield client
    app.dependency_overrides.clear()


def test_pages_and_validation(client):
    first = client.get("/api/internships?limit=10").json()
    second = client.get("/api/internships?limit=10&page=2").json()
    assert first["total"] == 15
    assert len(first["items"]) == 10 and len(second["items"]) == 5
    assert not {j["id"] for j in first["items"]} & {j["id"] for j in second["items"]}
    assert client.get("/api/internships?page=0").status_code == 422
    assert client.get("/api/internships?sort=stipend").status_code == 422
    assert client.get("/api/internships?role=backend").json()["total"] == 0
    assert client.get("/api/internships?min_match=101").status_code == 422


def test_tracking_corrections_and_exports(client):
    job = client.get("/api/internships").json()["items"][0]
    endpoint = "/api/applications/"+job["id"]
    updated = client.patch(endpoint,json={"application_status":"applied","notes":"Follow up","favorite":True}).json()
    assert updated["applied_at"]
    assert client.patch(endpoint,json={"notes":"Changed"}).json()["applied_at"] == updated["applied_at"]
    assert client.patch(endpoint,json={"application_status":"invalid"}).status_code == 422
    assert client.patch(endpoint,json={"notes":None}).status_code == 422
    assert client.get("/api/applications").json()[0]["notes"] == "Changed"
    corrected = client.patch("/api/internships/"+job["id"]+"/corrections",json={"remote_status":"onsite"}).json()
    assert corrected["remote_status"] == "onsite"
    assert client.get("/api/internships?favorite=true").json()["total"] == 1
    csv = client.get("/api/export.csv?favorite=true")
    assert csv.status_code == 200 and csv.text.count("Frontend Intern") == 1
    xlsx = client.get("/api/export.xlsx?favorite=true")
    assert load_workbook(BytesIO(xlsx.content)).active.max_row == 2


def test_profile_search_and_analytics(client):
    profile = CandidateProfile(skills=["Python"],preferred_roles=["backend"])
    assert client.put("/api/profile",json=profile.model_dump()).status_code == 200
    assert client.get("/api/profile").json()["skills"] == ["Python"]
    assert client.post("/api/searches",json={"name":"Remote paid","filters":{"remote":"remote","paid":"true"}}).status_code == 200
    assert len(client.get("/api/searches").json()) == 1
    assert client.delete("/api/searches/Remote%20paid").json() == []
    assert client.get("/api/analytics/market").json()["total"] == 15
    assert len(client.get("/api/providers/health").json()) == 8


def test_resume_limits_privacy_and_security(client):
    result = client.post("/api/resume/analyze",content=b"B.Tech university\nReact JavaScript",headers={"Content-Type":"text/plain"}).json()
    assert result["stored"] is False
    assert "react" in result["skills"]
    assert client.post("/api/resume/analyze",content=b"bad",headers={"Content-Type":"application/pdf"}).status_code == 422
    assert client.post("/api/resume/analyze",content=b"x"*2_000_001,headers={"Content-Type":"text/plain"}).status_code == 413
    assert client.get("/api/profile",headers={"Origin":"https://evil.example"}).status_code == 403
    pdf = PdfWriter()
    for _ in range(11):
        pdf.add_blank_page(width=200,height=200)
    out = BytesIO(); pdf.write(out)
    assert client.post("/api/resume/analyze",content=out.getvalue(),headers={"Content-Type":"application/pdf"}).status_code == 422


def test_export_formula_injection():
    job = make_job(title="=HYPERLINK(\"https://evil.example\")")
    assert "'=HYPERLINK" in jobs_to_csv([job])
    sheet = load_workbook(BytesIO(jobs_to_xlsx([job]))).active
    assert all(cell.data_type != "f" for row in sheet for cell in row)


def test_private_token_and_host_boundary(client, monkeypatch):
    from app.main import settings
    monkeypatch.setattr(settings,"api_token","test-only-secret")
    assert client.get("/api/profile").status_code == 401
    assert client.get("/api/profile",headers={"Authorization":"Bearer wrong"}).status_code == 401
    assert client.get("/api/profile",headers={"Authorization":"Bearer test-only-secret"}).status_code == 200
    assert client.get("/api/profile",headers={"Authorization":"Bearer test-only-secret","Host":"attacker.example"}).status_code == 400
    assert client.get("/health").status_code == 200


def test_saved_search_invalid_filters_rejected(client):
    assert client.post("/api/searches",json={"name":"Invalid","filters":{"min_match":"1000"}}).status_code == 422


def test_readiness_search_explanation_and_reset(client):
    assert client.get('/ready').json()['status'] == 'ready'
    data=client.get('/api/internships?q=paid%20frontend%20internships%20underwater').json()
    assert data['total'] == 0 and data['unparsed_query'] == 'underwater'
    assert data['search_warnings']
    job=client.get('/api/internships').json()['items'][0]
    path='/api/internships/'+job['id']+'/corrections'
    client.patch(path,json={'remote_status':'onsite'})
    assert client.patch(path,json={'reset':True}).json()['remote_status'] == 'remote'


@pytest.mark.parametrize('mime,content,status', [('image/png',b'png',415),('text/plain',b'\xff',422),('application/pdf',b'%PDF-invalid',422),('text/plain',b' ',422)])
def test_invalid_resume_uploads(client,mime,content,status):
    assert client.post('/api/resume/analyze',content=content,headers={'Content-Type':mime}).status_code == status


def test_quarantine_excluded_from_all_normal_routes(client):
    repo=app.dependency_overrides[get_repository]()
    asyncio.run(repo.upsert_jobs([make_job(url='https://boards.greenhouse.io/acme/jobs/scam',description='Pay registration fee before joining. Telegram only.')]))
    for path in ('/api/internships','/api/jobs','/api/recommendations'):
        data=client.get(path).json()
        jobs=data.get('items',[]) if isinstance(data,dict) else data
        assert all(j['risk_state'] not in ('quarantined','blocked') for j in jobs)
    assert len(client.get('/api/listings/risky').json()) == 1


def test_operator_enablement_visible_in_health(client,monkeypatch):
    from app.main import settings
    monkeypatch.setattr(settings,'enabled_sources',['yc_jobs'])
    providers=client.get('/api/providers/health').json()
    assert all(p['status']=='disabled' for p in providers if p['source']!='yc_jobs')
